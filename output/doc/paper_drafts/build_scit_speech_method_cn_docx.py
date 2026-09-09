from __future__ import annotations

import json
import math
import os
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
SOURCE_MD = ROOT / "output/doc/paper_drafts/scit_speech_method_cn_draft_20260622.md"
OUTPUT_DOCX = Path(os.environ.get("SCIT_DOCX_OUTPUT", SOURCE_MD.with_suffix(".docx"))).resolve()

BODY_EAST_ASIA = "SimSun"
TITLE_EAST_ASIA = "SimHei"
CAPTION_EAST_ASIA = "FangSong"
LATIN_FONT = "Times New Roman"
MATH_FONT = "Times New Roman"

NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
nsmap.setdefault("m", NS_M)


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] | None = None
    aligns: list[str] | None = None
    alt: str = ""
    path: str = ""
    ordered: bool = False


def oxml(tag: str) -> OxmlElement:
    return OxmlElement(tag)


def set_rfonts(element, ascii_font=LATIN_FONT, east_asia=BODY_EAST_ASIA, hansi=None):
    rpr = element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = oxml("w:rFonts")
        rpr.append(rfonts)
    hansi = hansi or ascii_font
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), hansi)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), ascii_font)


def set_run_font(run, size: float | None = None, bold=None, italic=None,
                 ascii_font=LATIN_FONT, east_asia=BODY_EAST_ASIA,
                 color: str | None = None):
    set_rfonts(run._element, ascii_font=ascii_font, east_asia=east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size: float, ascii_font=LATIN_FONT, east_asia=BODY_EAST_ASIA,
                   bold=None, color: str | None = None):
    font = style.font
    font.name = ascii_font
    font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color:
        font.color.rgb = RGBColor.from_string(color)
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = oxml("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), ascii_font)


def set_spacing(paragraph, before=0, after=6, line=1.25):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_section_a4(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width = Mm(297)
        section.page_height = Mm(210)
    else:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.header_distance = Mm(12.5)
    section.footer_distance = Mm(12.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, east_asia=CAPTION_EAST_ASIA)
    fld_begin = oxml("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = oxml("w:instrText")
    instr.set(XML_SPACE, "preserve")
    instr.text = " PAGE "
    fld_sep = oxml("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = oxml("w:t")
    fld_text.text = "1"
    fld_end = oxml("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = oxml("w:r")
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(fld_text)
    r.append(fld_end)
    paragraph._p.append(r)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, east_asia=CAPTION_EAST_ASIA)


def configure_document(doc: Document):
    for section in doc.sections:
        set_section_a4(section, landscape=False)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.first_line_indent = Pt(21)

    for name, size, color, before, after in [
        ("Heading 1", 14, "1F4D78", 14, 7),
        ("Heading 2", 12.5, "1F4D78", 10, 5),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        set_style_font(style, size, east_asia=TITLE_EAST_ASIA, bold=True, color=color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.2
        style.paragraph_format.first_line_indent = Pt(0)

    if "Caption" in styles:
        cap = styles["Caption"]
        set_style_font(cap, 9.5, east_asia=CAPTION_EAST_ASIA, bold=False, color="222222")
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(8)
        cap.paragraph_format.line_spacing = 1.15
        cap.paragraph_format.first_line_indent = Pt(0)

    for style_name in ["List Paragraph", "List Number", "List Bullet"]:
        if style_name in styles:
            style = styles[style_name]
            set_style_font(style, 10.5)
            style.paragraph_format.left_indent = Pt(21)
            style.paragraph_format.first_line_indent = Pt(-21)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.2

    section = doc.sections[0]
    header_p = section.header.paragraphs[0]
    header_p.text = ""
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrun = header_p.add_run("SCIT-Speech 中文论文稿")
    set_run_font(hrun, size=9, east_asia=CAPTION_EAST_ASIA, color="666666")

    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    add_page_number(footer_p)


def is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\|\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line))


def is_table_line(line: str) -> bool:
    return line.strip().startswith("|") and "|" in line.strip()[1:]


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def parse_aligns(sep: str) -> list[str]:
    aligns = []
    for cell in parse_table_row(sep):
        left = cell.startswith(":")
        right = cell.endswith(":")
        if left and right:
            aligns.append("center")
        elif right:
            aligns.append("right")
        else:
            aligns.append("left")
    return aligns


def starts_block(lines: list[str], i: int) -> bool:
    line = lines[i]
    stripped = line.strip()
    if not stripped:
        return True
    if stripped.startswith("```"):
        return True
    if re.match(r"^#{1,4}\s+", stripped):
        return True
    if re.match(r"^!\[[^\]]*\]\([^)]+\)\s*$", stripped):
        return True
    if i + 1 < len(lines) and is_table_line(line) and is_table_sep(lines[i + 1]):
        return True
    if re.match(r"^\s*(?:[-*]\s+|\d+\.\s+)", line):
        return True
    return False


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append(Block(kind="math_block", text="\n".join(buf).strip("\n")))
            continue
        hm = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if hm:
            blocks.append(Block(kind="heading", level=len(hm.group(1)), text=hm.group(2).strip()))
            i += 1
            continue
        im = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", stripped)
        if im:
            blocks.append(Block(kind="image", alt=im.group(1), path=im.group(2).strip()))
            i += 1
            continue
        if i + 1 < len(lines) and is_table_line(line) and is_table_sep(lines[i + 1]):
            rows = [parse_table_row(line)]
            aligns = parse_aligns(lines[i + 1])
            i += 2
            while i < len(lines) and is_table_line(lines[i]):
                rows.append(parse_table_row(lines[i]))
                i += 1
            blocks.append(Block(kind="table", rows=rows, aligns=aligns))
            continue
        lm = re.match(r"^\s*(?:(\d+)\.\s+|[-*]\s+)(.*)$", line)
        if lm:
            ordered = lm.group(1) is not None
            text_part = lm.group(2).strip()
            i += 1
            while i < len(lines) and lines[i].strip() and not starts_block(lines, i):
                text_part += " " + lines[i].strip()
                i += 1
            blocks.append(Block(kind="list", text=text_part, ordered=ordered))
            continue
        buf = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not starts_block(lines, i):
            buf.append(lines[i].strip())
            i += 1
        blocks.append(Block(kind="paragraph", text=" ".join(buf)))
    return blocks


def is_caption_text(text: str) -> bool:
    clean = re.sub(r"^\*\*(.*?)\*\*", r"\1", text.strip())
    return bool(re.match(r"^(图|表)\s*[\dD][\w.\-a-zA-Z]*", clean))


def strip_outer_bold(text: str) -> tuple[str, bool]:
    m = re.match(r"^\*\*(.*?)\*\*(.*)$", text.strip())
    if m:
        return m.group(1) + m.group(2), True
    return text, False


def split_bold_definition(text: str) -> tuple[str | None, str, str]:
    m = re.match(r"^\*\*([^*\n]{1,80})\*\*\s*([:\uFF1A])\s*(.*)$", text.strip())
    if not m:
        return None, "", text.strip()
    return m.group(1).strip(), m.group(2), m.group(3).strip()


def math_text_to_omath(text: str):
    omath = OxmlElement("m:oMath")
    for child in math_text_children(normalize_math_text(text)):
        omath.append(child)
    return omath


def math_run(text: str):
    mr = OxmlElement("m:r")
    mrpr = OxmlElement("m:rPr")
    sty = OxmlElement("m:sty")
    sty.set(qn("m:val"), "p")
    mrpr.append(sty)
    mr.append(mrpr)
    wrpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), MATH_FONT)
    rfonts.set(qn("w:hAnsi"), MATH_FONT)
    rfonts.set(qn("w:eastAsia"), MATH_FONT)
    rfonts.set(qn("w:cs"), MATH_FONT)
    wrpr.append(rfonts)
    mr.append(wrpr)
    mt = OxmlElement("m:t")
    mt.set(XML_SPACE, "preserve")
    mt.text = text
    mr.append(mt)
    return mr


def math_script(base: str, sub: str | None = None, sup: str | None = None):
    if sub and sup:
        node = OxmlElement("m:sSubSup")
        sub_tag = "m:sub"
        sup_tag = "m:sup"
    elif sub:
        node = OxmlElement("m:sSub")
        sub_tag = "m:sub"
        sup_tag = None
    else:
        node = OxmlElement("m:sSup")
        sub_tag = None
        sup_tag = "m:sup"

    e = OxmlElement("m:e")
    e.append(math_run(base))
    node.append(e)
    if sub_tag:
        sub_node = OxmlElement(sub_tag)
        sub_node.append(math_run(sub or ""))
        node.append(sub_node)
    if sup_tag:
        sup_node = OxmlElement(sup_tag)
        sup_node.append(math_run(sup or ""))
        node.append(sup_node)
    return node


def consume_script(text: str, pos: int) -> tuple[str, int]:
    if pos >= len(text):
        return "", pos
    if text[pos] == "{":
        end = text.find("}", pos + 1)
        if end != -1:
            return text[pos + 1:end], end + 1
    if text[pos] == "(":
        end = text.find(")", pos + 1)
        if end != -1:
            return text[pos:end + 1], end + 1
    j = pos
    while j < len(text) and re.match(r"[A-Za-z0-9βΔλφΣ:,]+", text[j]):
        j += 1
    return text[pos:j], j


def parse_scripted_token(token: str) -> tuple[str, str | None, str | None]:
    m = re.match(r"([A-Za-zλφΔΣβĨ][A-Za-z0-9λφΔΣβĨ̂]*)", token)
    if not m:
        return token, None, None
    base = m.group(1)
    pos = len(base)
    sub_parts: list[str] = []
    sup_parts: list[str] = []
    while pos < len(token):
        mark = token[pos]
        if mark not in "_^":
            break
        pos += 1
        value, pos = consume_script(token, pos)
        if mark == "_":
            sub_parts.append(value)
        else:
            sup_parts.append(value)
    sub = ",".join(part for part in sub_parts if part)
    sup = ",".join(part for part in sup_parts if part)
    return base, sub or None, sup or None


SCRIPTED_TOKEN_RE = re.compile(
    r"[A-Za-zλφΔΣβĨ][A-Za-z0-9λφΔΣβĨ̂]*(?:(?:_\{[^}]+\}|_[A-Za-z0-9βΔλφΣ:,]+|\^\{[^}]+\}|\^[A-Za-z0-9βΔλφΣ()]+))+"
)


def math_text_children(text: str):
    children = []
    pos = 0
    for match in SCRIPTED_TOKEN_RE.finditer(text):
        if match.start() > pos:
            children.append(math_run(text[pos:match.start()]))
        base, sub, sup = parse_scripted_token(match.group(0))
        children.append(math_script(base, sub=sub, sup=sup))
        pos = match.end()
    if pos < len(text):
        children.append(math_run(text[pos:]))
    return children


def normalize_math_text(text: str) -> str:
    replacements = {
        "theta": "θ",
        "lambda": "λ",
        "phi": "φ",
        "Delta": "Δ",
        "Sigma": "Σ",
        "\\lambda": "λ",
        "\\theta": "θ",
        "\\phi": "φ",
        "\\Delta": "Δ",
        "\\Sigma": "Σ",
        "\\in": "∈",
        "\\le": "≤",
        "\\ge": "≥",
        "\\neq": "≠",
        "\\cdot": "·",
        "\\times": "×",
        "\\hat{x}": "x̂",
        "\\hat{I}": "Ĩ",
        "\\lceil": "⌈",
        "\\rceil": "⌉",
        "\\log_2": "log₂",
    }
    out = text
    for old, new in replacements.items():
        out = out.replace(old, new)
    out = out.replace("...", "…")
    out = out.replace("（", "(").replace("）", ")")
    return out


def add_inline_math(paragraph, text: str):
    paragraph._p.append(math_text_to_omath(text))


def add_plain_run(paragraph, text: str, size: float, bold=False, italic=False,
                  east_asia=BODY_EAST_ASIA, ascii_font=LATIN_FONT,
                  color: str | None = None):
    if not text:
        return None
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic,
                 east_asia=east_asia, ascii_font=ascii_font, color=color)
    return run


def add_inline_markdown(paragraph, text: str, size: float = 10.5, bold_default=False,
                        east_asia=BODY_EAST_ASIA):
    i = 0
    while i < len(text):
        if text.startswith("**", i):
            j = text.find("**", i + 2)
            if j != -1:
                add_inline_markdown(paragraph, text[i + 2:j], size=size,
                                    bold_default=True, east_asia=east_asia)
                i = j + 2
                continue
        if text[i] == "`":
            j = text.find("`", i + 1)
            if j != -1:
                code = text[i + 1:j]
                add_inline_math(paragraph, code)
                i = j + 1
                continue
        add_plain_run(paragraph, text[i], size=size, bold=bold_default, east_asia=east_asia)
        i += 1


def add_paragraph(doc: Document, text: str, style=None, size=10.5, align=None,
                  first_line=True, east_asia=BODY_EAST_ASIA):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.first_line_indent = Pt(21) if first_line else Pt(0)
    set_spacing(p, after=6, line=1.25)
    if align is not None:
        p.alignment = align
    add_inline_markdown(p, text, size=size, east_asia=east_asia)
    return p


def split_equation_number(line: str) -> tuple[str, str]:
    m = re.search(r"\s*[\(（]\s*(NAS)\s*[−\-]\s*(\d+)\s*[\)）]\s*$", line)
    if not m:
        return line.rstrip(), ""
    return line[:m.start()].rstrip(), f"({m.group(1)}-{m.group(2)})"


def split_formula_label(line: str) -> tuple[str, str]:
    m = re.match(r"^([^:=：]{1,12}[:：])\s*(.+)$", line.strip())
    if not m:
        return "", line.strip()
    return m.group(1), m.group(2).strip()


def set_no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = oxml("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = oxml(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def add_formula_cell_paragraph(cell, text: str, size=10.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = align
    p._p.append(math_text_to_omath(text))
    return p


def add_label_cell_paragraph(cell, text: str, align=WD_ALIGN_PARAGRAPH.RIGHT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=10.0, east_asia=CAPTION_EAST_ASIA, ascii_font=LATIN_FONT)
    return p


def add_display_math(doc: Document, text: str):
    rows = []
    for raw_line in text.splitlines() or [text]:
        line = raw_line.strip()
        if not line:
            continue
        formula, number = split_equation_number(line)
        label, formula = split_formula_label(formula)
        rows.append((label, formula, number))
    if not rows:
        return

    has_label = any(label for label, _, _ in rows)
    label_w = 1500 if has_label else 450
    num_w = 1200 if any(number for _, _, number in rows) else 450
    formula_w = 9070 - label_w - num_w
    table = doc.add_table(rows=len(rows), cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_no_table_borders(table)
    set_cell_margins(table, top=30, start=30, bottom=30, end=30)
    set_table_widths(table, [label_w, formula_w, num_w])
    for ri, (label, formula, number) in enumerate(rows):
        cells = table.rows[ri].cells
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        add_label_cell_paragraph(cells[0], label)
        add_formula_cell_paragraph(cells[1], formula)
        add_label_cell_paragraph(cells[2], number, align=WD_ALIGN_PARAGRAPH.RIGHT)

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_after = Pt(4)


def resolve_image(raw_path: str) -> Path:
    return (SOURCE_MD.parent / raw_path).resolve()


def add_image(doc: Document, raw_path: str, alt: str, missing: list[str]):
    path = resolve_image(raw_path)
    if not path.exists():
        missing.append(f"{raw_path} -> {path}")
        add_paragraph(doc, f"[缺失图片：{raw_path}]", size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                      first_line=False, east_asia=CAPTION_EAST_ASIA)
        return
    with Image.open(path) as im:
        w_px, h_px = im.size
    max_width_cm = 15.8
    width_cm = min(max_width_cm, max(8.0, max_width_cm if w_px >= h_px else 12.8))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    inline = run.add_picture(str(path), width=Cm(width_cm))
    doc_pr = inline._inline.docPr
    doc_pr.set("name", alt or path.name)
    doc_pr.set("descr", alt or path.name)


def set_cell_text(cell, text: str, size: float, bold=False, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline_markdown(p, text, size=size, bold_default=bold)


def set_table_borders(table):
    """Apply a standard academic three-line table style.

    Data tables get only top, header-bottom, and bottom rules. Formula layout
    tables call set_no_table_borders() instead and are intentionally unaffected.
    """
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = oxml("w:tblBorders")
        tbl_pr.append(borders)
    for edge, size in [("top", "12"), ("bottom", "12")]:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = oxml(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
    for edge in ["left", "right", "insideH", "insideV"]:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = oxml(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "nil")


def set_cell_border(cell, edge: str, val="single", size="8", color="000000"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = oxml("w:tcBorders")
        tc_pr.append(borders)
    el = borders.find(qn(f"w:{edge}"))
    if el is None:
        el = oxml(f"w:{edge}")
        borders.append(el)
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), size)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)


def set_header_bottom_rule(row):
    for cell in row.cells:
        set_cell_border(cell, "bottom", size="8")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(oxml("w:tblHeader"))


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = oxml("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=90, bottom=80, end=90):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = oxml("w:tblCellMar")
        tbl_pr.append(margins)
    for m_name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = margins.find(qn(f"w:{m_name}"))
        if node is None:
            node = oxml(f"w:{m_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa: list[int]):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = oxml("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = oxml("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        col = oxml("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = oxml("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def content_weight(cell: str) -> float:
    ascii_chars = sum(1 for c in cell if ord(c) < 128)
    non_ascii = len(cell) - ascii_chars
    return ascii_chars * 0.65 + non_ascii * 1.0


def compute_widths(rows: list[list[str]], total_dxa: int) -> list[int]:
    cols = max(len(r) for r in rows)
    weights = []
    for c in range(cols):
        max_w = max(content_weight(r[c]) if c < len(r) else 1 for r in rows[: min(len(rows), 8)])
        weights.append(max(4.0, min(max_w, 28.0)))
    short_cols = cols >= 8
    min_w = 680 if short_cols else 900
    raw = [max(min_w, int(total_dxa * w / sum(weights))) for w in weights]
    scale = total_dxa / sum(raw)
    widths = [max(520, int(w * scale)) for w in raw]
    diff = total_dxa - sum(widths)
    widths[-1] += diff
    return widths


def add_table(doc: Document, rows: list[list[str]], aligns: list[str] | None, landscape=False):
    cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < cols:
            r.append("")
    total_dxa = 9070
    widths = compute_widths(rows, total_dxa)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_cell_margins(table)
    set_table_widths(table, widths)

    font_size = 7.2 if cols >= 10 else 8.0 if cols >= 8 else 8.8 if cols >= 6 else 9.2
    aligns = aligns or ["left"] * cols
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            align = aligns[ci] if ci < len(aligns) else "left"
            if ri == 0:
                align = "center"
            if ci > 0 and re.fullmatch(r"[+\-−]?\d[\d.,\s\[\]\-+<>/%×~]*", cell_text.replace(" ", "")):
                align = "center"
            set_cell_text(cell, cell_text, size=font_size, bold=(ri == 0), align=align)
    if table.rows:
        set_header_bottom_rule(table.rows[0])
        set_repeat_table_header(table.rows[0])
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.first_line_indent = Pt(0)
    return table


def add_heading(doc: Document, level: int, text: str):
    if level == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(14)
        run = p.add_run(text)
        set_run_font(run, size=18, bold=True, east_asia=TITLE_EAST_ASIA, color="000000")
        return p
    style = "Heading 1" if level == 2 else "Heading 2" if level == 3 else "Heading 3"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.first_line_indent = Pt(0)
    add_inline_markdown(p, text, size=14 if level == 2 else 12.5 if level == 3 else 11.5,
                        bold_default=True, east_asia=TITLE_EAST_ASIA)
    return p


def add_caption(doc: Document, text: str):
    clean, _ = strip_outer_bold(text)
    p = doc.add_paragraph(style="Caption" if "Caption" in doc.styles else None)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)
    add_inline_markdown(p, clean, size=9.5, east_asia=CAPTION_EAST_ASIA)
    return p


def is_definition_item_text(text: str) -> bool:
    """Detect bold-leading definition/explanation items such as **label**: text."""
    stripped = text.strip()
    if is_caption_text(stripped):
        return False
    label, _, _ = split_bold_definition(stripped)
    if not label:
        return False
    return label not in {"\u5173\u952e\u8bcd"}


def add_parenthesized_item(doc: Document, text: str, number: int, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.first_line_indent = Pt(21)
    set_spacing(p, after=6, line=1.25)
    add_plain_run(p, f"\uFF08{number}\uFF09", size=size)
    add_inline_markdown(p, text, size=size)
    return p


def add_list_item(doc: Document, text: str, ordered: bool):
    p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
    p.paragraph_format.left_indent = Pt(21)
    p.paragraph_format.first_line_indent = Pt(-21)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    add_inline_markdown(p, text, size=10.5)
    return p


def start_section(doc: Document, landscape: bool):
    section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_section_a4(section, landscape=landscape)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    return section


def add_reference_paragraph(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.first_line_indent = Pt(-24)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_inline_markdown(p, text, size=9.2)
    return p


def is_reference_line(text: str) -> bool:
    return bool(re.match(r"^\[\d+\]\s+", text.strip()))


def build_docx() -> dict:
    text = SOURCE_MD.read_text(encoding="utf-8")
    blocks = parse_markdown(text)

    doc = Document()
    configure_document(doc)
    missing_images: list[str] = []
    item_counter = 1
    current_heading = ""

    i = 0
    while i < len(blocks):
        block = blocks[i]
        next_block = blocks[i + 1] if i + 1 < len(blocks) else None

        if block.kind == "paragraph" and is_caption_text(block.text) and next_block and next_block.kind == "table":
            add_caption(doc, block.text)
            add_table(doc, next_block.rows, next_block.aligns, landscape=False)
            i += 2
            continue

        if block.kind == "heading":
            current_heading = block.text.strip()
            add_heading(doc, block.level, block.text)
            item_counter = 1
        elif block.kind == "paragraph":
            if is_caption_text(block.text):
                add_caption(doc, block.text)
            elif is_reference_line(block.text):
                add_reference_paragraph(doc, block.text)
            elif current_heading == "\u0034.4 \u8bc4\u4ef7\u6307\u6807":
                label, _, rest = split_bold_definition(block.text)
                if label == "\u5ba2\u89c2\u6307\u6807":
                    add_paragraph(doc, block.text)
                    item_counter = 1
                elif label == "\u4efb\u52a1\u5c42\u6307\u6807":
                    add_parenthesized_item(doc, block.text, item_counter)
                    item_counter += 1
                elif label == "WER \u6d4b\u91cf\u5730\u677f\uff08PCM \u4e0b\u754c\uff09":
                    add_paragraph(doc, f"**{label}**：")
                    if rest:
                        add_paragraph(doc, rest)
                    item_counter = 1
                elif is_definition_item_text(block.text):
                    add_parenthesized_item(doc, block.text, item_counter)
                    item_counter += 1
                else:
                    add_paragraph(doc, block.text)
            elif is_definition_item_text(block.text):
                add_parenthesized_item(doc, block.text, item_counter)
                item_counter += 1
            else:
                add_paragraph(doc, block.text)
        elif block.kind == "list":
            add_parenthesized_item(doc, block.text, item_counter)
            item_counter += 1
        elif block.kind == "math_block":
            add_display_math(doc, block.text)
        elif block.kind == "image":
            add_image(doc, block.path, block.alt, missing_images)
        elif block.kind == "table":
            add_table(doc, block.rows, block.aligns, landscape=False)
        i += 1

    doc.save(OUTPUT_DOCX)
    return verify_docx(missing_images)


def docx_text(doc: Document) -> str:
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def count_xml(pattern: str, xml: str) -> int:
    return len(re.findall(pattern, xml))


def verify_docx(missing_images: list[str]) -> dict:
    exists = OUTPUT_DOCX.exists() and OUTPUT_DOCX.stat().st_size > 0
    doc = Document(OUTPUT_DOCX)
    text = docx_text(doc)

    with zipfile.ZipFile(OUTPUT_DOCX) as zf:
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
        media = [n for n in zf.namelist() if n.startswith("word/media/")]

    required = {
        "title": "SCIT-Speech：基于共享 RVQ 码本索引传输的极低码率语音通信" in text,
        "abstract": "摘要" in text,
        "keywords": "关键词" in text,
        "section_1": "1 引言" in text,
        "section_8": "8 结论" in text,
        "references": "参考文献" in text and "[30]" in text,
        "appendix_d": "附录 D" in text,
        "figure_1_caption": "图 1" in text,
        "figure_6_caption": "图 6" in text,
    }
    limitation_items = 0
    in_limitations = False
    for p in doc.paragraphs:
        if p.text.strip() == "7 局限性":
            in_limitations = True
            continue
        if in_limitations and p.text.strip() == "8 结论":
            break
        if in_limitations and (
            p.style.name == "List Number" or re.match(r"^\uFF08[1-4]\uFF09", p.text.strip())
        ):
            limitation_items += 1
    raw_markdown_leaks = {
        "image_syntax": "![" in text,
        "pipe_table_lines": any(re.match(r"^\|.*\|$", line) for line in text.splitlines()),
        "many_backticks": text.count("`") > 20,
    }
    math_count = count_xml(r"<m:oMath", document_xml)
    return {
        "output": str(OUTPUT_DOCX),
        "exists": exists,
        "size_bytes": OUTPUT_DOCX.stat().st_size if OUTPUT_DOCX.exists() else 0,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(media),
        "math_omml_objects": math_count,
        "missing_images": missing_images,
        "required_checks": required,
        "references_1_to_30_present": all(f"[{i}]" in text for i in range(1, 31)),
        "limitation_items_in_section_7": limitation_items,
        "framework_word_count": text.lower().count("framework"),
        "markdown_leaks": raw_markdown_leaks,
        "all_required_present": all(required.values()),
        "no_markdown_leaks": not any(raw_markdown_leaks.values()),
    }


if __name__ == "__main__":
    result = build_docx()
    print(json.dumps(result, ensure_ascii=False, indent=2))
