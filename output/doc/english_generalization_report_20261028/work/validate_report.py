import csv
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path

import fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
EXP = ROOT.parents[1] / 'experiments' / 'english_generalization_3corpora_20261028_seed42'
checks = {}

for label, relative in [('primary', 'full'), ('control', 'level_control/full')]:
    folder = EXP / relative
    asr = defaultdict(lambda: {'n': 0, 'words': 0, 'errors': 0})
    objective = defaultdict(list)
    seen_asr, seen_obj = set(), set()
    for line in (folder / 'asr.jsonl').open(encoding='utf-8'):
        row = json.loads(line)
        uid = row['dataset'], row['sample_id']
        assert uid not in seen_asr
        seen_asr.add(uid)
        for model, result in row.get('asr', {}).items():
            if 'errors' not in result or 'reference_words' not in result:
                continue
            group = asr[(row['dataset'], model)]
            group['n'] += 1
            group['words'] += result['reference_words']
            group['errors'] += result['errors']
    for expected in csv.DictReader((folder / 'reports/asr_summary.csv').open(encoding='utf-8-sig')):
        raw = asr[(expected['dataset'], expected['model'])]
        assert raw['n'] == int(expected['n'])
        assert raw['words'] == int(expected['reference_words'])
        assert raw['errors'] == int(expected['errors'])
        assert math.isclose(raw['errors'] / raw['words'], float(expected['estimate']), abs_tol=1e-12)
    for line in (folder / 'objective.jsonl').open(encoding='utf-8'):
        row = json.loads(line)
        uid = row['dataset'], row['sample_id']
        assert uid not in seen_obj
        seen_obj.add(uid)
        for result in row['metrics']:
            for metric in ['mel_l1', 'stoi', 'pesq_wb', 'si_snr_db']:
                value = result.get(metric)
                if value is not None and math.isfinite(value):
                    objective[(row['dataset'], result['model'], str(result['L']), metric)].append(value)
    for expected in csv.DictReader((folder / 'reports/objective_summary.csv').open(encoding='utf-8-sig')):
        values = objective[(expected['dataset'], expected['model'], expected['L'], expected['metric'])]
        assert len(values) == int(expected['n'])
        assert math.isclose(sum(values) / len(values), float(expected['estimate']), abs_tol=1e-10)
    checks[label] = {'unique_asr_samples': len(seen_asr), 'unique_objective_samples': len(seen_obj),
                     'asr_groups': len(asr), 'objective_metric_groups': len(objective),
                     'aggregate_recalculation': 'passed'}

blocks = json.loads((ROOT / 'work/report_blocks.json').read_text(encoding='utf-8'))
document = Document(ROOT / '英文跨语料泛化评测报告.docx')
table_blocks = [b for b in blocks if b['type'] == 'table']
assert len(document.tables) == len(table_blocks)
for table, block in zip(document.tables, table_blocks):
    assert [[c.text for c in row.cells] for row in table.rows] == [block['headers']] + block['rows']

main_blocks = blocks[:next(i for i, b in enumerate(blocks) if b.get('text') == '附录 FLEURS 输入增益对照')]
main_text = json.dumps(main_blocks, ensure_ascii=False)
assert '低输入电平' not in main_text
assert '13.98' not in main_text
assert '57.17' in main_text
assert '5.55' in main_text and '9.28' in main_text
for b in blocks:
    for relative in b.get('paths', []):
        assert (EXP / relative).is_file(), relative

render_dir = ROOT / 'work/render'
pdf = fitz.open(render_dir / '英文跨语料泛化评测报告.pdf')
page_info = []
for index, page in enumerate(pdf):
    text = page.get_text()
    png_path = render_dir / f'page-{index+1}.png'
    page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False).save(png_path)
    page_info.append({'page': index+1, 'characters': len(text),
                      'first_lines': text.splitlines()[:4], 'last_lines': text.splitlines()[-4:]})
all_pdf_text = '\n'.join(page.get_text() for page in pdf)
for expected in ['57.17', '5.55', '9.28', '13.98', '2,200', '1,155', '647']:
    assert expected in all_pdf_text, expected
checks['document'] = {'tables': len(document.tables), 'inline_images': len(document.inline_shapes),
                      'page_count': len(pdf), 'table_content_check': 'passed',
                      'main_and_appendix_separation': 'passed', 'source_paths': 'passed',
                      'render_method': 'WPS native PDF export then PyMuPDF page rasterization',
                      'packaged_renderer_diagnostic': 'Bundled LibreOffice is unavailable on Windows; soffice.exe was not found on PATH',
                      'pages': page_info}
checks['deliverables'] = {p.name: {'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
                           for p in ROOT.iterdir() if p.suffix in ['.docx', '.md']}
(ROOT / 'work/validation.json').write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(checks, ensure_ascii=False, indent=2))
