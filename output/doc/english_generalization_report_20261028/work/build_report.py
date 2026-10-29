from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
EXP = ROOT.parents[1] / 'experiments' / 'english_generalization_3corpora_20261028_seed42'
WORK = ROOT / 'work'
FIG = ROOT / 'figures'
for folder in (ROOT, WORK, FIG):
    folder.mkdir(parents=True, exist_ok=True)


def read_csv(relative):
    with (EXP / relative).open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle))


asr_rows = read_csv('full/reports/asr_summary.csv')
asr_delta_rows = read_csv('full/reports/asr_paired_deltas.csv')
obj_rows = read_csv('full/reports/objective_summary.csv')
obj_delta_rows = read_csv('full/reports/objective_paired_deltas.csv')
ctrl_asr_rows = read_csv('level_control/full/reports/asr_summary.csv')
ctrl_obj_rows = read_csv('level_control/full/reports/objective_summary.csv')
manifest = json.loads((EXP / 'manifests/summary.json').read_text(encoding='utf-8'))
identity = json.loads((EXP / 'full/objective_identity.json').read_text(encoding='utf-8'))
validation = json.loads((EXP / 'checks/final_validation.json').read_text(encoding='utf-8'))

external = ['vctk', 'tedlium', 'fleurs_en_us']
reference = ['librispeech_test_clean', 'librispeech_test_other']
all_sets = external + reference
names = {
    'vctk': 'VCTK', 'tedlium': 'TED-LIUM', 'fleurs_en_us': 'FLEURS en_us',
    'librispeech_test_clean': 'LS test-clean', 'librispeech_test_other': 'LS test-other',
}


def find(rows, **keys):
    matches = [r for r in rows if all(str(r[k]) == str(v) for k, v in keys.items())]
    assert len(matches) == 1, (keys, len(matches))
    return matches[0]


def asr(ds, model, rows=asr_rows):
    return find(rows, dataset=ds, model=model)


def objective(ds, model, metric, layer=3, rows=obj_rows):
    return find(rows, dataset=ds, model=model, metric=metric, L=layer)


def pct(row, key='estimate'):
    return f'{100 * float(row[key]):.2f}'


def ci(row, scale=100, signed=False):
    spec = '+.2f' if signed else '.2f'
    return f'[{format(scale * float(row["ci_low"]), spec)}, {format(scale * float(row["ci_high"]), spec)}]'


blocks = []


def p(text, kind='body'):
    blocks.append({'type': kind, 'text': text})


def heading(text, page=False, level=1):
    blocks.append({'type': 'heading', 'text': text, 'page': page, 'level': level})


def table(caption, headers, rows, widths, font=10):
    blocks.append({'type': 'table', 'caption': caption, 'headers': headers,
                   'rows': rows, 'widths': widths, 'font': font})


def source(text, paths):
    blocks.append({'type': 'source', 'text': text, 'paths': paths})


p('SCIT Speech 英文跨语料泛化评测报告', 'title')
p('2026 年 10 月 28 日  |  三个外部英文语料与 LibriSpeech 参照', 'subtitle')
p('现有模型在 VCTK 和 TED-LIUM 上表现出一定的英文跨语料迁移能力。'
  '在 L3、1.5 kbit/s 条件下，LCA 的词错误率分别为 5.55% 和 9.28%。'
  'LCA 在 VCTK 上优于 Base，在 TED-LIUM 上基本持平。FLEURS 的主实验词错误率为 57.17%，'
  '说明当前模型的跨语料表现仍不均衡。', 'lead')
p('本轮实验已在实验室远端主机 lab-suer 的 suer 账户、物理 GPU 1 上完成。'
  '评测使用现有 Base 与 LCA 检查点，冻结全部权重，不在外部测试语料上训练、适配或选择模型。'
  '报告重点回答模型能否将已有英文语音编码能力迁移到外部语料，以及 LCA 的收益能否保持。')

heading('评测范围', level=2)
scope_rows = []
for ds in all_sets:
    m = manifest['datasets'][ds]
    cluster = {'speaker': '说话人', 'talk': '演讲', 'prompt': '文本 prompt'}[m['cluster_type']]
    scope_rows.append([names[ds], f'{m["count"]:,}', f'{int(asr(ds,"lca")["n"]):,}',
                       f'{m["clusters"]} {cluster}', f'{m["duration_hours"]:.3f}'])
table('表 1  数据覆盖与统计单元',
      ['语料', '客观评测条数', 'WER 条数', '聚类单元', '时长 小时'],
      scope_rows, [1.52, 1.08, .86, 2.04, 1.40])
p('三个外部语料共 4,002 条录音，另有 LibriSpeech 两个测试子集各 300 条作为参照。'
  '主实验合计 4,602 条录音、27,612 组模型与码率组合记录，完成 13,746 次有效 ASR 识别。'
  'LS 为 LibriSpeech 的缩写。', 'note')
p('VCTK 采用 mic1，每位说话人抽取 20 条，共 110 位说话人；其中包含此前已测的 300 条，'
  '新增 1,900 条。p315 的 20 条缺少参考转写，仅参与客观评测，因此 WER 的说话人聚类数为 109。')
p('TED-LIUM 使用固定版本的 release3 test 镜像。1,469 条源记录中，314 条按 '
  'ignore_time_segment_in_scoring 标记排除，保留 1,155 条。FLEURS 使用 en_us test 的全部 647 条录音。')
source('数据依据', ['PROTOCOL.md', 'manifests/summary.json', 'full/reports/coverage.csv'])

heading('评测设置与主实验结果', page=True)
p('模型与码率。Base 使用 SCIT-Speech-Base_best.pt，LCA 使用 '
  'SCIT-Speech-LCA_v2_step30000_robust_optimum.pt。'
  '模型按原配置严格加载，使用 FP32 推理。一次编码后分别解码 L1、L2、L3 前缀，'
  '名义载荷码率为 0.5、1.0、1.5 kbit/s，不含传输开销。L3 为主要工作点，L1/L2 为补充检查；本轮未注入丢包。')
p('统一预处理。整句音频转换为 float32 单声道并重采样至 16 kHz，保留原始幅值。'
  '不额外裁剪、按时长筛选或进行时间对齐。客观指标在参考与重建音频的共同长度上计算，'
  '并记录两者长度。ASR 使用完整原音与共同长度的 L3 重建音频。')
p('识别设置。固定使用 faster-whisper-small.en，英文模式、beam size 5、temperature 0，'
  '关闭 VAD 与前文条件，推理精度为 float16。参考文本与识别文本使用同一固定版本的 '
  'Whisper 英文归一化器处理。WER 由全语料替换、删除和插入词数之和除以参考词数计算，不取逐句 WER 的平均值。')
p('指标与不确定性。WER 和 mel-L1 越低越好，STOI、PESQ-WB、SI-SNR 越高越好。'
  '95% 区间由 2,000 次聚类 bootstrap 得到，随机种子为 42。VCTK 与 LibriSpeech 按说话人、'
  'TED-LIUM 按演讲、FLEURS 按文本 prompt 重采样；模型差异采用配对重采样。')

wer_rows = []
for ds in external:
    wer_rows.append([names[ds], pct(asr(ds, 'original')), pct(asr(ds, 'base')),
                     pct(asr(ds, 'lca')), ci(asr(ds, 'lca'))])
table('表 2  三个外部英文语料的 L3 词错误率',
      ['语料', '原音 WER %', 'Base WER %', 'LCA WER %', 'LCA 的 95% 区间 %'],
      wer_rows, [1.38, 1.03, 1.03, 1.03, 2.43])
p('原音 WER 是同一识别器对未经过编码重建的录音测得的错误率，用于判断编码带来的额外误差。'
  'LCA 相对原音的 WER 增量，在 VCTK、TED-LIUM、FLEURS 上分别为 3.87、4.89、49.87 个百分点。'
  '这些值应与原音基线共同解读。')
p('VCTK 的 LCA 结果接近 LibriSpeech test-clean 参照的量级；TED-LIUM 中大部分内容仍可被固定识别器识别。'
  'FLEURS 的绝对错误率较高，构成本轮跨语料泛化的明显边界。')
source('结果依据', ['full/reports/asr_summary.csv', 'full/reports/asr_paired_deltas.csv', 'PROTOCOL.md'])

heading('模型比较与 LibriSpeech 参照', page=True)
shutil.copy2(EXP / 'figures/english_generalization_wer_L3.png', FIG / 'english_generalization_wer_L3.png')
blocks.append({'type': 'image', 'path': str(FIG / 'english_generalization_wer_L3.png'),
               'width': 6.88,
               'caption': '图 1  L3 的语料级 WER 与相对原音的 WER 增量。误差棒为 95% 聚类 bootstrap 区间。'
                          'LS clean 与 LS other 为 LibriSpeech 参照。该图反映自动识别结果，不是人工听测评分。'})
paired_rows = []
for ds in external:
    r = find(asr_delta_rows, dataset=ds, contrast='lca-base')
    paired_rows.append([names[ds], f'{100*float(r["estimate"]):+.2f}', ci(r, signed=True)])
table('表 3  LCA 相对 Base 的配对 WER 差异',
      ['语料', 'LCA 减 Base 百分点', '配对 95% 区间 百分点'],
      paired_rows, [1.55, 2.35, 3.00])
p('负值表示 LCA 的 WER 更低。VCTK 的差异为 −0.91 个百分点，区间不跨 0，'
  '支持该语料和工作点上的改善。TED-LIUM 的差异仅为 +0.01 个百分点，区间跨 0，'
  '尚无明确的 WER 优势证据。FLEURS 虽降低 3.42 个百分点，两模型的绝对错误率仍然较高。')
ref_rows = []
for ds in reference:
    ref_rows.append([names[ds], pct(asr(ds, 'original')), pct(asr(ds, 'base')),
                    pct(asr(ds, 'lca')), ci(asr(ds, 'lca'))])
table('表 4  LibriSpeech 参照结果',
      ['语料', '原音 WER %', 'Base WER %', 'LCA WER %', 'LCA 的 95% 区间 %'],
      ref_rows, [1.38, 1.03, 1.03, 1.03, 2.43])
p('test-clean 上 LCA 比 Base 降低 0.82 个百分点；test-other 上 LCA 的点估计高 0.55 个百分点，'
  '其配对 95% 区间为 [−1.46, +2.86]，不足以确定差异方向。'
  '不同语料的原音难度、内容和录音条件同时变化，因此不将跨语料分数差直接解释为单一因素的作用。')
source('结果依据', ['full/reports/asr_summary.csv', 'full/reports/asr_paired_deltas.csv', 'figures/CAPTIONS.md'])

heading('客观音质与码率检查', page=True)
metric_order = ['mel_l1', 'stoi', 'pesq_wb', 'si_snr_db']
objective_table = []
for ds in all_sets:
    for model in ['base', 'lca']:
        objective_table.append([names[ds], model.upper() if model == 'lca' else 'Base'] +
                               [f'{float(objective(ds, model, metric)["estimate"]):.4f}' for metric in metric_order])
table('表 5  L3 客观指标均值',
      ['语料', '模型', 'mel-L1 ↓', 'STOI ↑', 'PESQ-WB ↑', 'SI-SNR dB ↑'],
      objective_table, [1.42, .60, 1.14, 1.04, 1.26, 1.44], font=9.5)
p('VCTK 与 TED-LIUM 上，LCA 的 mel-L1 和 STOI 优于 Base，PESQ-WB 与 SI-SNR 的均值略低。'
  'FLEURS 上，LCA 的四项客观指标均弱于 Base，但 WER 较低。'
  '各指标反映的性能侧面不同，不能将单项改善等同于音质全面提升。')

layer_rows = []
for ds in external:
    layer_rows.append([names[ds]] + [f'{float(find(obj_delta_rows,dataset=ds,L=l,metric="mel_l1")["estimate"]):+.4f}'
                                    for l in [1, 2, 3]])
table('表 6  不同码率的 mel-L1 配对差值',
      ['语料', 'L1  0.5 kbit/s', 'L2  1.0 kbit/s', 'L3  1.5 kbit/s'],
      layer_rows, [1.50, 1.80, 1.80, 1.80])
p('表 6 均为 LCA 减 Base，负值表示 LCA 更好。VCTK 在 L1/L2 上出现 mel-L1 退化，'
  '到 L3 才改善；TED-LIUM 三个工作点均改善；FLEURS 三个工作点均未改善。'
  '因此，L3 的优势不能外推到所有码率。本轮 L1/L2 未运行 ASR，不能据此推断其 WER。')
p('TED-LIUM 的 ted_00769 在两模型、三个工作点下均无法计算 PESQ，共 6 个失败条目。'
  '各 TED-LIUM PESQ 分组的有效样本数为 1,154，其他指标为 1,155。'
  '失败项保留在记录中，未补零。所有具有参考文本的 ASR 样本均已完成。', 'note')
source('结果依据', ['full/reports/objective_summary.csv', 'full/reports/objective_paired_deltas.csv',
                   'full/reports/metric_failures.csv'])

heading('结论与证据范围', page=True)
p('本轮完成了三个外部英文语料的冻结模型评测。结果支持模型在 VCTK 和 TED-LIUM 上具有一定的跨语料迁移能力，'
  '其中 LCA 在 VCTK 的 L3 词错误率上有明确改善，TED-LIUM 的词错误率与 Base 基本持平。'
  'FLEURS 的主实验结果明显退化，说明泛化表现依赖具体语料条件。', 'lead')
p('LCA 的收益具有指标和码率边界。其在 VCTK、TED-LIUM 的 L3 mel-L1 与 STOI 上取得改善，'
  '但 PESQ-WB、SI-SNR 并未同步提升；VCTK 的低码率 mel-L1 也未保持 L3 的优势。'
  '当前证据适合支持“部分外部英文语料与工作点上的迁移及改进”，尚不足以支持“所有英文语料或码率下全面领先”。')
heading('统计与数据边界', level=2)
p('统计区间来自同一对现有训练种子的检查点，反映评测样本变化，不包含重复训练造成的波动。'
  '多语料、多指标比较未进行多重比较校正。TED-LIUM 仅有 11 个演讲聚类；FLEURS 缺少说话人标识，'
  '采用 350 个文本 prompt 聚类，不能据此声称控制了说话人相关性。')
p('已检查 Base 与 LCA 的训练、验证清单：各含 27,039 条训练记录、1,500 条验证记录，'
  '未发现与本轮测试重叠的样本 ID，也未发现三个外部语料的路径记录。'
  '该检查范围限于现有清单，不构成对全部预训练数据的完整去污染证明。ASR 与其他预训练组件的语料暴露情况尚未建立。')
p('WER 来自固定自动识别器，客观音质指标也不能替代人工听测。本轮未引入丢包或开展新的主观质量评估，'
  '相关通信鲁棒性与感知质量结论需由对应实验支持。所有跨语料比较均使用本轮一致的运行环境，未混入历史评测分数。')
heading('结果追溯', level=2)
p('实验根目录为 output/experiments/english_generalization_3corpora_20261028_seed42。'
  '完整远端记录保存在 /home/suer/english_generalization_20261028。以下路径均相对实验根目录。')
p('协议与样本：PROTOCOL.md、RUNBOOK.md、manifests/all_samples.jsonl。', 'source')
p('主结果：full/objective.jsonl、full/asr.jsonl，以及 full/reports 下的统计汇总、配对差异、覆盖率和失败记录。', 'source')
p('复现证据：provenance/training_manifest_audit.json、provenance/source_hashes.json、'
  'full/objective_identity.json、full/asr_identity.json。', 'source')
p('完成核验：checks/final_validation.json 记录主实验与补充对照的聚合结果重算通过。'
  '本报告中的表格直接从归档 CSV 生成，并按逐条识别记录与客观结果校验。', 'source')

heading('附录 FLEURS 输入增益对照', page=True)
p('此项为观察主实验结果后设计的补充对照，覆盖 FLEURS 全部 647 条录音。'
  '它采用独立预处理条件，结果与正文主实验分别报告，不替代主实验估计，也不属于预先设计的独立确认实验。')
p('对每条录音施加恒定增益，目标 RMS 为 −25 dBFS，峰值上限为 0.99。'
  '6 条录音因峰值限制未达到目标 RMS。模型、参考文本与评分方式保持一致，无权重更新、去噪或按效果筛除样本。')
control_table = []
for label, ar, ob in [('主实验', asr_rows, obj_rows), ('增益对照', ctrl_asr_rows, ctrl_obj_rows)]:
    control_table.append([label, pct(asr('fleurs_en_us', 'original', ar)),
                          pct(asr('fleurs_en_us', 'base', ar)), pct(asr('fleurs_en_us', 'lca', ar)),
                          f'{float(objective("fleurs_en_us","lca","stoi",rows=ob)["estimate"]):.4f}'])
table('表 A1  FLEURS 补充对照的 L3 结果',
      ['条件', '原音 WER %', 'Base WER %', 'LCA WER %', 'LCA STOI'],
      control_table, [1.30, 1.40, 1.40, 1.40, 1.40])
p('增益对照下，LCA 的 WER 为 13.98%，95% 区间为 [12.73%, 15.45%]。'
  '与主实验同样本配对比较，WER 降低 43.19 个百分点，95% 区间为 [40.05, 46.41]。'
  '原音 WER 从 7.30% 变为 7.25%，基本不变；LCA 相对原音仍增加 6.73 个百分点。')
p('对应 LCA 客观指标为 mel-L1 0.8268、STOI 0.8141、PESQ-WB 1.8162、SI-SNR 1.5627 dB。'
  '这组结果表明评测表现会随输入处理条件变化，其作用范围限于本次补充对照。')
source('补充对照依据', ['level_control/CONTROL_PROTOCOL.json', 'level_control/full/reports/asr_summary.csv',
                      'level_control/full/reports/objective_summary.csv', 'analysis/fleurs_level_control_paired.json'])
heading('检查点与运行标识', level=2)
p('Base 检查点 SHA256', 'source')
p(identity['models']['base']['checkpoint_sha256'], 'hash')
p('LCA 检查点 SHA256', 'source')
p(identity['models']['lca']['checkpoint_sha256'], 'hash')
p('主实验样本清单 SHA256', 'source')
p(manifest['manifest_sha256'], 'hash')
p('运行记录：Python 3.10.20，PyTorch 2.10.0，torchaudio 2.10.0+cu128，'
  'faster-whisper 1.2.1。数据版本、ASR 模型版本、文本归一化器版本、完整依赖与源代码哈希均见原始实验归档。', 'note')


# A single content representation drives both editable Word and Markdown.
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.8)
sec.right_margin = Inches(.8)
sec.footer_distance = Inches(.28)


def font_setup(rpr, latin='Calibri', east='Microsoft YaHei'):
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for attr in list(rf.attrib):
        if 'theme' in attr.lower():
            del rf.attrib[attr]
    for name, value in [('ascii', latin), ('hAnsi', latin), ('eastAsia', east), ('cs', latin)]:
        rf.set(qn('w:' + name), value)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rpr.append(lang)
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:eastAsia'), 'zh-CN')


for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.italic = False
    st.font.underline = False
    font_setup(st.element.get_or_add_rPr())
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = Pt(17)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True

for name, size, line, before, after in [
    ('Title', 24, 31, 0, 9), ('Subtitle', 10.5, 16, 0, 14),
    ('Heading 1', 18, 25, 0, 11), ('Heading 2', 12.5, 19, 9, 5),
    ('Caption', 9, 13.5, 3, 6),
]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = name in ['Title', 'Heading 1', 'Heading 2']
    st.paragraph_format.line_spacing = Pt(line)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

for name, size, line in [('Note', 9.5, 14.5), ('Source', 9, 13), ('Hash', 8.5, 12.5)]:
    st = doc.styles.add_style(name, 1)
    st.base_style = normal
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.line_spacing = Pt(line)
    st.paragraph_format.space_after = Pt(5)
    font_setup(st.element.get_or_add_rPr(), latin='Consolas' if name == 'Hash' else 'Calibri')


def add_paragraph(text, style='Normal', bold=False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    return paragraph


def add_table(block):
    cap = add_paragraph(block['caption'], 'Caption')
    cap.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(block['headers']))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, block['widths']):
        col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement('w:' + name)
        for attr, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
            border.set(qn('w:' + attr), value)
        borders.append(border)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for name, value in [('top', '76'), ('bottom', '76'), ('left', '85'), ('right', '85')]:
        margin = OxmlElement('w:' + name)
        margin.set(qn('w:w'), value)
        margin.set(qn('w:type'), 'dxa')
        margins.append(margin)
    props.append(margins)
    rows = [block['headers']] + block['rows']
    for i, values in enumerate(rows):
        row = t.rows[0] if i == 0 else t.add_row()
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(block['widths'][j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '203C50' if i == 0 else ('F0F5F8' if i % 2 == 0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(shade)
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            pf = par.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = Pt(14)
            pf.keep_with_next = i < len(rows) - 1
            run = par.add_run(str(value))
            run.font.size = Pt(block['font'])
            run.font.bold = i == 0
            run.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '000000')
            font_setup(run._r.get_or_add_rPr())
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.add_run().font.size = Pt(2)


md = []
for block in blocks:
    kind = block['type']
    if kind == 'heading':
        par = add_paragraph(block['text'], f'Heading {block["level"]}')
        par.paragraph_format.page_break_before = block['page']
        md += ['#' * (block['level'] + 1) + ' ' + block['text'], '']
    elif kind == 'table':
        add_table(block)
        md += [block['caption'], '', '| ' + ' | '.join(block['headers']) + ' |',
               '| ' + ' | '.join(['---'] * len(block['headers'])) + ' |']
        md += ['| ' + ' | '.join(str(x) for x in row) + ' |' for row in block['rows']]
        md += ['']
    elif kind == 'image':
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.line_spacing = 1
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.keep_with_next = True
        image = par.add_run().add_picture(block['path'], width=Inches(block['width']))
        image._inline.docPr.set('descr', block['caption'])
        add_paragraph(block['caption'], 'Caption')
        image_relative = Path(os.path.relpath(block['path'], ROOT)).as_posix()
        md += [f'![L3 英文跨语料 WER 与相对原音的差异]({image_relative})', '', block['caption'], '']
    elif kind == 'source' and 'paths' in block:
        add_paragraph(block['text'] + '：' + '；'.join(block['paths']) + '。', 'Source')
        links = [f'[{name}]({Path(os.path.relpath(EXP / name, ROOT)).as_posix()})' for name in block['paths']]
        md += [block['text'] + '：' + '；'.join(links) + '。', '']
    else:
        style = {'title': 'Title', 'subtitle': 'Subtitle', 'source': 'Source', 'note': 'Note', 'hash': 'Hash'}.get(kind, 'Normal')
        add_paragraph(block['text'], style, bold=kind == 'lead')
        text = block['text']
        if kind == 'title':
            text = '# ' + text
        elif kind == 'lead':
            text = '**' + text + '**'
        elif kind == 'hash':
            text = '`' + text + '`'
        md += [text, '']

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_after = Pt(0)
for text in ['第 ', 'PAGE', ' 页']:
    run = footer.add_run()
    run.font.size = Pt(9)
    if text == 'PAGE':
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), 'PAGE')
        run._r.addnext(fld)
    else:
        run.text = text

doc.core_properties.title = 'SCIT Speech 英文跨语料泛化评测报告'
doc.core_properties.subject = 'VCTK TED-LIUM FLEURS 与 LibriSpeech 的冻结模型评测'
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.keywords = 'SCIT-Speech, English, generalization, frozen evaluation'
docx_path = ROOT / '英文跨语料泛化评测报告.docx'
md_path = ROOT / '英文跨语料泛化评测报告.md'
doc.save(docx_path)
md_path.write_text('\n'.join(md).strip() + '\n', encoding='utf-8')
(WORK / 'report_blocks.json').write_text(json.dumps(blocks, ensure_ascii=False, indent=2), encoding='utf-8')
used_sources = ['PROTOCOL.md', 'manifests/summary.json', 'full/reports/asr_summary.csv',
                'full/reports/asr_paired_deltas.csv', 'full/reports/objective_summary.csv',
                'full/reports/objective_paired_deltas.csv', 'full/reports/metric_failures.csv',
                'provenance/training_manifest_audit.json', 'checks/final_validation.json',
                'full/objective_identity.json', 'level_control/CONTROL_PROTOCOL.json',
                'level_control/full/reports/asr_summary.csv',
                'level_control/full/reports/objective_summary.csv',
                'analysis/fleurs_level_control_paired.json', 'figures/english_generalization_wer_L3.png']
hashes = {s: hashlib.sha256((EXP / s).read_bytes()).hexdigest() for s in used_sources}
(WORK / 'source_sha256.json').write_text(json.dumps(hashes, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'docx': str(docx_path), 'markdown': str(md_path),
                  'planned_pages': 6, 'tables': len([b for b in blocks if b['type'] == 'table']),
                  'source_count': len(hashes)}, ensure_ascii=False))
